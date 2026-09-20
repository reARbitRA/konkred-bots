"""Deep Document Assistant - summaries, quizzes, flashcards and risk analysis.

Documents are downloaded into memory, text is extracted locally (pypdf /
python-docx / plain decode) and the resulting corpus is routed to the gateway
with ``taskType="spec-generation"``, which prioritises Gemini Flash and its 1M
token context window so entire contracts fit in a single call.
"""

from __future__ import annotations

import io
import json
import logging

from aiogram import Bot, F, Router
from aiogram.enums import ChatAction
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.types import CallbackQuery, Message

from shared.config import settings
from shared.gateway_client import GatewayError, gateway
from shared.history import HistoryManager
from shared.payments import PaymentManager
from shared.utils import (
    UNEXPECTED_ERROR,
    clean_model_output,
    escape_html,
    extract_json_block,
    humanize_bytes,
    send_long_message,
    truncate,
)

from .keyboards import (
    CB_PREFIX,
    document_actions,
    finish_keyboard,
    next_question_keyboard,
    quiz_answer_keyboard,
)

logger = logging.getLogger(__name__)
router = Router(name="bot_pdf")

# Documents are truncated to this many characters before being sent upstream.
# ~600k chars ≈ 170k tokens, comfortably inside every long-context candidate.
MAX_DOC_CHARS = 600_000

SYSTEM_PROMPT = (
    "You are Konkred DocMind, a meticulous document analyst. You ground every statement in the "
    "supplied document, quote figures exactly, and explicitly say when information is absent "
    "rather than guessing."
)

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx", ".csv", ".json", ".log", ".rst"}

WELCOME = """📄 <b>Deep Document Assistant</b>

Send me a document and I'll read all of it — then you choose what to do with it.

<b>I accept</b>: PDF · DOCX · TXT · MD · CSV · JSON (up to {max_mb} MB)

<b>I can produce</b>
📝 Executive summary
❓ An interactive 5-question quiz
📇 Flashcards for revision
⚠️ Contract risk analysis
🔑 Key terms · 📊 Data & figures

You can also just ask me questions about the document in plain language.

<b>Commands</b>: /start · /help · /clear"""

HELP = """📄 <b>Deep Document Assistant — help</b>

<b>1. Send a document</b>
PDF, DOCX, TXT, MD, CSV or JSON, up to {max_mb} MB.
I extract the text locally, then analyse it with a 1M-token context model, so
long contracts and reports are processed whole rather than in fragments.

<b>2. Pick an action</b>
• <b>Executive Summary</b> — purpose, key points, conclusions
• <b>Quiz</b> — 5 multiple-choice questions, asked one at a time with scoring
• <b>Flashcards</b> — Q/A pairs for spaced repetition
• <b>Contract Risk Analysis</b> — liabilities, obligations, missing clauses, red flags
• <b>Key Terms</b> — glossary of domain vocabulary
• <b>Data & Figures</b> — every number, date and metric, in a table

<b>3. Or just ask</b>
"What is the termination notice period?" — I answer from the document only.

<b>Commands</b>: /start · /help · /clear (forget the current document)"""


# --------------------------------------------------------------------------- #
# Local text extraction
# --------------------------------------------------------------------------- #

def _extract_pdf(payload: bytes) -> tuple[str, str]:
    """Extract text from a PDF, tolerating individually malformed pages."""
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - dependency is pinned
        return "", "pypdf is not installed"

    try:
        reader = PdfReader(io.BytesIO(payload))
    except Exception as exc:  # noqa: BLE001 - pypdf raises many parse error types
        logger.warning("pdf parse failed: %s", exc)
        return "", f"the PDF could not be parsed ({exc})"

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:  # noqa: BLE001
            return "", "the PDF is password-protected"

    pages: list[str] = []
    try:
        total = len(reader.pages)
    except Exception as exc:  # noqa: BLE001
        return "", f"the PDF could not be read ({exc})"

    for number, page in enumerate(reader.pages, start=1):
        try:
            content = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001 - one bad page must not fail the file
            logger.warning("pdf page %s extraction failed: %s", number, exc)
            continue
        if content.strip():
            pages.append(f"--- Page {number} ---\n{content.strip()}")

    text = "\n\n".join(pages)
    if not text.strip():
        return "", "this PDF contains no extractable text (it is probably a scan)"
    return text, f"{total} page(s)"


def _extract_docx(payload: bytes) -> tuple[str, str]:
    """Extract paragraphs and tables from a DOCX file."""
    try:
        import docx
    except ImportError:  # pragma: no cover - dependency is pinned
        return "", "python-docx is not installed"

    try:
        document = docx.Document(io.BytesIO(payload))
    except Exception as exc:  # noqa: BLE001
        logger.warning("docx parse failed: %s", exc)
        return "", f"the DOCX file could not be parsed ({exc})"

    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table_index, table in enumerate(document.tables, start=1):
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
            if any(cell.text.strip() for cell in row.cells)
        ]
        if rows:
            blocks.append(f"--- Table {table_index} ---\n" + "\n".join(rows))

    text = "\n\n".join(blocks)
    if not text.strip():
        return "", "this DOCX file contains no text"
    return text, f"{len(document.paragraphs)} paragraph(s)"


def _extract_plain(payload: bytes) -> tuple[str, str]:
    """Decode a plain-text family file, trying several encodings."""
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = payload.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
        if text.strip():
            return text, f"{len(text.splitlines())} line(s)"
        return "", "the file is empty"
    return "", "the file is not readable as text"


def extract_text(payload: bytes, filename: str, mime_type: str = "") -> tuple[str, str]:
    """Extract plain text from a document.

    Returns:
        ``(text, note)`` where ``note`` either describes the extraction
        (e.g. "12 page(s)") or, when ``text`` is empty, explains the failure in
        language suitable for showing directly to the user.
    """
    lower = (filename or "").lower()
    mime = mime_type or ""

    if lower.endswith(".pdf") or mime == "application/pdf":
        return _extract_pdf(payload)
    if lower.endswith(".docx") or "wordprocessingml" in mime:
        return _extract_docx(payload)
    return _extract_plain(payload)


# --------------------------------------------------------------------------- #
# FSM-free document state (kept in aiogram's Redis-backed FSM storage)
# --------------------------------------------------------------------------- #

async def _store_document(state: FSMContext, text: str, filename: str, note: str) -> None:
    await state.update_data(
        doc_text=text[:MAX_DOC_CHARS],
        doc_name=filename,
        doc_note=note,
        doc_chars=len(text),
        quiz=None,
        quiz_score=0,
        quiz_answered=0,
    )


async def _document(state: FSMContext) -> tuple[str, str]:
    data = await state.get_data()
    return data.get("doc_text", ""), data.get("doc_name", "document")


def _doc_messages(doc_text: str, doc_name: str, instruction: str) -> list[dict]:
    return [{
        "role": "user",
        "content": (
            f"<document name=\"{doc_name}\">\n{doc_text}\n</document>\n\n{instruction}"
        ),
    }]


# --------------------------------------------------------------------------- #
# Commands
# --------------------------------------------------------------------------- #

@router.message(CommandStart())
async def cmd_start(message: Message) -> None:
    await message.answer(WELCOME.format(max_mb=settings.max_file_mb))


@router.message(Command("help"))
async def cmd_help(message: Message) -> None:
    await message.answer(HELP.format(max_mb=settings.max_file_mb))


@router.message(Command("clear"))
async def cmd_clear(message: Message, state: FSMContext, history: HistoryManager) -> None:
    await state.clear()
    await history.clear(message.from_user.id)
    await message.answer("🧹 Cleared. Send me a new document whenever you're ready.")


# --------------------------------------------------------------------------- #
# Document ingestion
# --------------------------------------------------------------------------- #

@router.message(F.document)
async def handle_document(message: Message, bot: Bot, state: FSMContext) -> None:
    document = message.document
    filename = document.file_name or "document"
    lower = filename.lower()

    if not any(lower.endswith(extension) for extension in SUPPORTED_EXTENSIONS):
        await message.answer(
            f"📎 I can't read <b>{escape_html(filename)}</b>.\n\n"
            f"Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
        return

    if (document.file_size or 0) > settings.max_file_bytes:
        await message.answer(
            f"📦 That file is {humanize_bytes(document.file_size)}, over my "
            f"{settings.max_file_mb} MB limit."
        )
        return

    status = await message.answer(f"📥 Reading <b>{escape_html(filename)}</b>…")
    await bot.send_chat_action(message.chat.id, ChatAction.TYPING)

    try:
        file = await bot.get_file(document.file_id)
        buffer = await bot.download_file(file.file_path)
        payload = buffer.read()
    except Exception as exc:  # noqa: BLE001
        logger.exception("document download failed")
        await status.edit_text(f"⚠️ I couldn't download that file.\n<code>{exc}</code>")
        return

    text, note = extract_text(payload, filename, document.mime_type or "")
    if not text.strip():
        await status.edit_text(f"⚠️ I couldn't extract any text: {note}.")
        return

    truncated = len(text) > MAX_DOC_CHARS
    await _store_document(state, text, filename, note)

    logger.info("pdf: user=%s file=%s chars=%s", message.from_user.id, filename, len(text))

    await status.edit_text(
        f"✅ <b>{escape_html(filename)}</b> loaded\n"
        f"📊 {note} · {len(text):,} characters"
        + ("\n⚠️ <i>Very long document — analysing the first 600,000 characters.</i>" if truncated else "")
        + "\n\nWhat would you like me to do?",
        reply_markup=document_actions(),
    )


# --------------------------------------------------------------------------- #
# Analysis actions
# --------------------------------------------------------------------------- #

ANALYSIS_PROMPTS = {
    "summary": (
        "Write an executive summary of this document with these sections:\n"
        "**Purpose** — what the document is for, in one sentence.\n"
        "**Key Points** — 5 to 8 bullets covering the substance.\n"
        "**Conclusions / Outcomes** — what it decides, recommends or requires.\n"
        "**Who should care** — the audiences affected and why.\n"
        "Quote figures and dates exactly as they appear."
    ),
    "flashcards": (
        "Create 10 revision flashcards from this document.\n"
        "Format each exactly as:\n\n"
        "**Card N**\n🔵 <b>Q:</b> question\n🟢 <b>A:</b> answer\n\n"
        "Cover the most examinable facts, definitions and relationships. "
        "Answers must be self-contained and under 40 words."
    ),
    "risk": (
        "Perform a contract/commercial risk analysis of this document. Use these sections:\n"
        "**🔴 High Risk** — clauses that create serious exposure; quote them.\n"
        "**🟡 Medium Risk** — unfavourable or ambiguous terms.\n"
        "**🟢 Standard** — ordinary, market-standard provisions.\n"
        "**❓ Missing Protections** — clauses a prudent party would expect but which are absent "
        "(liability caps, termination rights, IP assignment, confidentiality, dispute resolution, indemnities).\n"
        "**📋 Obligations Summary** — who must do what, by when.\n"
        "For each risk state the clause reference, the exposure it creates and a concrete mitigation. "
        "End with: 'This is an automated analysis, not legal advice.' "
        "If the document is not a contract, analyse its commercial and operational risks instead."
    ),
    "terms": (
        "Extract a glossary of the key terms, acronyms and domain vocabulary in this document. "
        "Format: **Term** — plain-language definition grounded in how the document uses it. "
        "Order by importance, 10-20 entries."
    ),
    "data": (
        "Extract every quantitative fact from this document: figures, amounts, percentages, dates, "
        "deadlines, durations and metrics. Present them grouped by theme as lines of "
        "`• Label: value (context / where it appears)`. Do not round or reinterpret any number. "
        "If the document contains no quantitative data, say so explicitly."
    ),
}


@router.callback_query(F.data == f"{CB_PREFIX}:menu")
async def callback_menu(query: CallbackQuery) -> None:
    await query.answer()
    await query.message.answer("What would you like me to do?", reply_markup=document_actions())


@router.callback_query(F.data.in_({f"{CB_PREFIX}:{action}" for action in ANALYSIS_PROMPTS}))
async def callback_analysis(
    query: CallbackQuery,
    state: FSMContext,
    history: HistoryManager,
    payments: PaymentManager,
) -> None:
    action = query.data.split(":", 1)[1]
    doc_text, doc_name = await _document(state)
    await query.answer()

    if not doc_text:
        await query.message.answer("📄 I don't have a document loaded. Please send one first.")
        return
    if not await payments.require(query.message, query.from_user.id):
        return

    labels = {
        "summary": "📝 Writing the executive summary…",
        "flashcards": "📇 Building flashcards…",
        "risk": "⚠️ Running the risk analysis…",
        "terms": "🔑 Extracting key terms…",
        "data": "📊 Pulling out the figures…",
    }
    status = await query.message.answer(labels.get(action, "⏳ Working…"))

    try:
        answer = await gateway.ask(
            task_type="spec-generation",
            messages=_doc_messages(doc_text, doc_name, ANALYSIS_PROMPTS[action]),
            system=SYSTEM_PROMPT,
            max_tokens=4000,
            temperature=0.25,
            user_id=query.from_user.id,
        )
    except GatewayError as exc:
        await status.edit_text(exc.user_message())
        return
    except Exception:  # noqa: BLE001 - last resort: never strand the user
        logger.exception("unexpected handler failure")
        await status.edit_text(UNEXPECTED_ERROR)
        return

    await history.add_exchange(query.from_user.id, f"[{action} of {doc_name}]", answer)
    await status.delete()
    await send_long_message(query.message, clean_model_output(answer))
    await query.message.answer("Anything else?", reply_markup=document_actions())


# --------------------------------------------------------------------------- #
# Interactive quiz
# --------------------------------------------------------------------------- #

QUIZ_PROMPT = """Create exactly 5 multiple-choice questions that test genuine comprehension of this document.

Return ONLY valid JSON, no prose and no code fence, in exactly this shape:
{
  "questions": [
    {
      "question": "the question text",
      "options": ["option A", "option B", "option C", "option D"],
      "correct": 0,
      "explanation": "why that option is correct, citing the document"
    }
  ]
}

Rules:
- exactly 5 questions, each with exactly 4 options
- "correct" is the 0-based index of the right option
- vary the position of the correct answer across questions
- every question must be answerable from the document alone
- no true/false questions; make the distractors plausible"""


@router.callback_query(F.data == f"{CB_PREFIX}:quiz")
async def callback_quiz(
    query: CallbackQuery,
    state: FSMContext,
    payments: PaymentManager,
) -> None:
    doc_text, doc_name = await _document(state)
    await query.answer()
    if not doc_text:
        await query.message.answer("📄 I don't have a document loaded. Please send one first.")
        return
    if not await payments.require(query.message, query.from_user.id):
        return

    status = await query.message.answer("❓ Writing your quiz…")
    try:
        raw = await gateway.ask(
            task_type="spec-generation",
            messages=_doc_messages(doc_text, doc_name, QUIZ_PROMPT),
            system=SYSTEM_PROMPT,
            max_tokens=3000,
            temperature=0.4,
            json_mode=True,
            user_id=query.from_user.id,
        )
    except GatewayError as exc:
        await status.edit_text(exc.user_message())
        return
    except Exception:  # noqa: BLE001 - last resort: never strand the user
        logger.exception("unexpected handler failure")
        await status.edit_text(UNEXPECTED_ERROR)
        return

    payload = extract_json_block(raw)
    questions = (payload or {}).get("questions") if isinstance(payload, dict) else None
    questions = [q for q in (questions or []) if _valid_question(q)]

    if not questions:
        logger.warning("quiz generation produced no valid questions: %s", truncate(raw, 300))
        await status.edit_text(
            "⚠️ I couldn't build a well-formed quiz from that document. "
            "Try the summary instead, or send a text-richer document.",
            reply_markup=document_actions(),
        )
        return

    await state.update_data(quiz=json.dumps(questions), quiz_score=0, quiz_answered=0)
    await status.delete()
    await _ask_question(query.message, state, 0)


def _valid_question(question: object) -> bool:
    """Structural validation so a malformed model response can never crash a handler."""
    if not isinstance(question, dict):
        return False
    options = question.get("options")
    correct = question.get("correct")
    return (
        isinstance(question.get("question"), str)
        and bool(question["question"].strip())
        and isinstance(options, list)
        and len(options) == 4
        and all(isinstance(option, str) and option.strip() for option in options)
        and isinstance(correct, int)
        and 0 <= correct < 4
    )


async def _ask_question(message: Message, state: FSMContext, index: int) -> None:
    data = await state.get_data()
    questions = json.loads(data.get("quiz") or "[]")
    if index >= len(questions):
        await _finish_quiz(message, state)
        return

    question = questions[index]
    letters = "ABCD"
    body = "\n".join(
        f"<b>{letters[i]}</b>) {escape_html(option)}"
        for i, option in enumerate(question["options"])
    )
    await message.answer(
        f"<b>Question {index + 1} of {len(questions)}</b>\n\n"
        f"{escape_html(question['question'])}\n\n{body}",
        reply_markup=quiz_answer_keyboard(index),
    )


@router.callback_query(F.data.startswith(f"{CB_PREFIX}:answer:"))
async def callback_answer(query: CallbackQuery, state: FSMContext) -> None:
    try:
        _, _, raw_index, raw_choice = query.data.split(":", 3)
        index, choice = int(raw_index), int(raw_choice)
    except (ValueError, IndexError):
        await query.answer("Invalid answer.", show_alert=True)
        return

    data = await state.get_data()
    questions = json.loads(data.get("quiz") or "[]")
    if index >= len(questions):
        await query.answer("That quiz has expired.", show_alert=True)
        return

    question = questions[index]
    correct = int(question["correct"])
    is_correct = choice == correct
    score = int(data.get("quiz_score", 0)) + (1 if is_correct else 0)
    answered = int(data.get("quiz_answered", 0)) + 1
    await state.update_data(quiz_score=score, quiz_answered=answered)

    await query.answer("✅ Correct!" if is_correct else "❌ Not quite.")

    letters = "ABCD"
    verdict = (
        f"✅ <b>Correct!</b> {letters[correct]}) {escape_html(question['options'][correct])}"
        if is_correct
        else (
            f"❌ You chose <b>{letters[choice]}</b>. "
            f"The answer is <b>{letters[correct]}</b>) {escape_html(question['options'][correct])}"
        )
    )
    explanation = question.get("explanation")
    if isinstance(explanation, str) and explanation.strip():
        verdict += f"\n\n💡 <i>{escape_html(explanation.strip())}</i>"
    verdict += f"\n\n<b>Score: {score}/{answered}</b>"

    try:
        await query.message.edit_reply_markup(reply_markup=None)
    except Exception:  # noqa: BLE001 - message may be too old to edit
        pass

    if index + 1 < len(questions):
        await query.message.answer(verdict, reply_markup=next_question_keyboard(index + 1))
    else:
        await query.message.answer(verdict)
        await _finish_quiz(query.message, state)


@router.callback_query(F.data.startswith(f"{CB_PREFIX}:next:"))
async def callback_next(query: CallbackQuery, state: FSMContext) -> None:
    await query.answer()
    try:
        index = int(query.data.rsplit(":", 1)[1])
    except ValueError:
        return
    await _ask_question(query.message, state, index)


@router.callback_query(F.data == f"{CB_PREFIX}:quizstop")
async def callback_quiz_stop(query: CallbackQuery, state: FSMContext) -> None:
    await query.answer("Quiz stopped.")
    await state.update_data(quiz=None, quiz_score=0, quiz_answered=0)
    await query.message.answer("🛑 Quiz stopped.", reply_markup=document_actions())


async def _finish_quiz(message: Message, state: FSMContext) -> None:
    data = await state.get_data()
    questions = json.loads(data.get("quiz") or "[]")
    total = len(questions) or 1
    score = int(data.get("quiz_score", 0))
    percent = round(score / total * 100)

    if percent >= 90:
        verdict = "🏆 Outstanding — you've mastered this document."
    elif percent >= 70:
        verdict = "👏 Strong result. A quick review of the misses and you're there."
    elif percent >= 50:
        verdict = "📚 Solid start — worth another pass over the key sections."
    else:
        verdict = "🔄 Worth re-reading. Try the flashcards, then retake the quiz."

    await state.update_data(quiz=None, quiz_score=0, quiz_answered=0)
    await message.answer(
        f"🎯 <b>Quiz complete</b>\n\n"
        f"Final score: <b>{score}/{total}</b> ({percent}%)\n\n{verdict}",
        reply_markup=finish_keyboard(),
    )


# --------------------------------------------------------------------------- #
# Free-form Q&A over the loaded document
# --------------------------------------------------------------------------- #

@router.message(F.text & ~F.text.startswith("/"))
async def handle_question(
    message: Message,
    state: FSMContext,
    history: HistoryManager,
    payments: PaymentManager,
) -> None:
    doc_text, doc_name = await _document(state)
    if not doc_text:
        await message.answer(
            "📄 Send me a PDF, DOCX, TXT or MD file first and I'll analyse it.\n\nUse /help for details."
        )
        return
    if not await payments.require(message, message.from_user.id):
        return

    thinking = await message.answer("🔍 Searching the document…")
    instruction = (
        f"Answer this question using ONLY the document above: {message.text}\n\n"
        "Quote the relevant passage. If the document does not contain the answer, say so plainly."
    )
    try:
        answer = await gateway.ask(
            task_type="spec-generation",
            messages=_doc_messages(doc_text, doc_name, instruction),
            system=SYSTEM_PROMPT,
            max_tokens=2500,
            temperature=0.2,
            user_id=message.from_user.id,
        )
    except GatewayError as exc:
        await thinking.edit_text(exc.user_message())
        return
    except Exception:  # noqa: BLE001 - last resort: never strand the user
        logger.exception("unexpected handler failure")
        await thinking.edit_text(UNEXPECTED_ERROR)
        return

    await history.add_exchange(message.from_user.id, message.text, answer)
    await thinking.delete()
    await send_long_message(message, clean_model_output(answer))


__all__ = ["router", "extract_text"]
